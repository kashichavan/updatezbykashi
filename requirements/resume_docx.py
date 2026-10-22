import io
import re
from typing import Dict, Any
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_horizontal_rule(paragraph):
    """Add a bottom border line to a heading paragraph in docx."""
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')  # 1pt border
    bottom.set(qn('w:space'), '2')
    bottom.set(qn('w:color'), '0F172A')
    pBdr.append(bottom)
    pPr.append(pBdr)


def generate_ats_resume_docx(resume_data: Dict[str, Any]) -> bytes:
    """
    Generates a true 100% native Microsoft Word .docx ATS-compliant resume.
    Properly formatted with margins, headers, bullet points, and typography.
    """
    doc = docx.Document()

    # Set A4 margins
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    # Base font Arial
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(15, 23, 42)

    # 1. Candidate Name (Header)
    name = (resume_data.get('name') or 'Candidate Name').strip().upper()
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run(name)
    run_name.bold = True
    run_name.font.size = Pt(18)
    run_name.font.name = 'Arial'
    run_name.font.color.rgb = RGBColor(15, 23, 42)

    # 2. Contact Information Line
    contact = resume_data.get('contact', {})
    contact_parts = []
    if contact.get('phone'):
        contact_parts.append(str(contact['phone']).strip())
    if contact.get('email'):
        contact_parts.append(str(contact['email']).strip())
    if contact.get('location'):
        contact_parts.append(str(contact['location']).strip())
    if contact.get('linkedin'):
        contact_parts.append(str(contact['linkedin']).strip())
    if contact.get('github'):
        contact_parts.append(str(contact['github']).strip())

    if contact_parts:
        p_contact = doc.add_paragraph()
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_contact.paragraph_format.space_before = Pt(0)
        p_contact.paragraph_format.space_after = Pt(8)
        run_contact = p_contact.add_run('  •  '.join(contact_parts))
        run_contact.font.size = Pt(9.5)
        run_contact.font.color.rgb = RGBColor(51, 65, 85)

    def add_section_header(title: str):
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(8)
        p_head.paragraph_format.space_after = Pt(3)
        p_head.paragraph_format.keep_with_next = True
        run_head = p_head.add_run(title.upper())
        run_head.bold = True
        run_head.font.size = Pt(11)
        run_head.font.color.rgb = RGBColor(15, 23, 42)
        add_horizontal_rule(p_head)
        return p_head

    # 3. Professional Summary
    summary = resume_data.get('summary')
    if summary:
        add_section_header("Professional Summary")
        p_sum = doc.add_paragraph()
        p_sum.paragraph_format.space_before = Pt(2)
        p_sum.paragraph_format.space_after = Pt(6)
        p_sum.paragraph_format.line_spacing = 1.15
        run_sum = p_sum.add_run(summary.strip())
        run_sum.font.size = Pt(9.5)

    # 4. Technical Skills
    skills = resume_data.get('skills', {})
    if skills:
        add_section_header("Technical Skills")
        if isinstance(skills, dict):
            for cat, items in skills.items():
                p_sk = doc.add_paragraph()
                p_sk.paragraph_format.space_before = Pt(1)
                p_sk.paragraph_format.space_after = Pt(2)
                r_cat = p_sk.add_run(f"{cat}: ")
                r_cat.bold = True
                r_cat.font.size = Pt(9.5)
                r_items = p_sk.add_run(str(items))
                r_items.font.size = Pt(9.5)
        elif isinstance(skills, list):
            p_sk = doc.add_paragraph()
            p_sk.paragraph_format.space_before = Pt(1)
            p_sk.paragraph_format.space_after = Pt(2)
            r_items = p_sk.add_run(', '.join([str(s) for s in skills]))
            r_items.font.size = Pt(9.5)

    # 5. Work Experience (if experienced or present)
    experience = resume_data.get('experience', [])
    if experience:
        add_section_header("Professional Experience")
        for exp in experience:
            p_exp = doc.add_paragraph()
            p_exp.paragraph_format.space_before = Pt(4)
            p_exp.paragraph_format.space_after = Pt(1)
            p_exp.paragraph_format.keep_with_next = True

            r_comp = p_exp.add_run(exp.get('company', 'Company'))
            r_comp.bold = True
            r_comp.font.size = Pt(10)

            if exp.get('period'):
                r_per = p_exp.add_run(f"  |  {exp.get('period')}")
                r_per.font.size = Pt(9.5)
                r_per.font.color.rgb = RGBColor(100, 116, 139)

            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(0)
            p_sub.paragraph_format.space_after = Pt(2)
            p_sub.paragraph_format.keep_with_next = True

            r_role = p_sub.add_run(exp.get('role', 'Software Engineer'))
            r_role.italic = True
            r_role.font.size = Pt(9.5)

            if exp.get('location'):
                r_loc = p_sub.add_run(f"  •  {exp.get('location')}")
                r_loc.font.size = Pt(9)
                r_loc.font.color.rgb = RGBColor(100, 116, 139)

            if exp.get('tech_stack'):
                p_tech = doc.add_paragraph()
                p_tech.paragraph_format.space_before = Pt(0)
                p_tech.paragraph_format.space_after = Pt(2)
                p_tech.paragraph_format.keep_with_next = True
                r_tl = p_tech.add_run("Core Stack: ")
                r_tl.bold = True
                r_tl.font.size = Pt(9)
                r_ts = p_tech.add_run(exp.get('tech_stack'))
                r_ts.font.size = Pt(9)

            for b in (exp.get('bullets') or []):
                p_b = doc.add_paragraph(style='List Bullet')
                p_b.paragraph_format.space_before = Pt(0)
                p_b.paragraph_format.space_after = Pt(1.5)
                r_b = p_b.add_run(b)
                r_b.font.size = Pt(9.5)

    # 6. Enterprise / Flagship Software Projects
    projects = resume_data.get('projects', [])
    if projects:
        add_section_header("Flagship Software Projects")
        for proj in projects:
            p_proj = doc.add_paragraph()
            p_proj.paragraph_format.space_before = Pt(4)
            p_proj.paragraph_format.space_after = Pt(1)
            p_proj.paragraph_format.keep_with_next = True

            r_title = p_proj.add_run(proj.get('title', 'Project Title'))
            r_title.bold = True
            r_title.font.size = Pt(10)

            if proj.get('role'):
                r_prole = p_proj.add_run(f"  |  {proj.get('role')}")
                r_prole.italic = True
                r_prole.font.size = Pt(9.5)
                r_prole.font.color.rgb = RGBColor(100, 116, 139)

            if proj.get('tech_stack'):
                p_ptech = doc.add_paragraph()
                p_ptech.paragraph_format.space_before = Pt(0)
                p_ptech.paragraph_format.space_after = Pt(2)
                p_ptech.paragraph_format.keep_with_next = True
                r_ptl = p_ptech.add_run("Technologies: ")
                r_ptl.bold = True
                r_ptl.font.size = Pt(9)
                r_pts = p_ptech.add_run(str(proj.get('tech_stack')))
                r_pts.font.size = Pt(9)

            for b in (proj.get('bullets') or []):
                p_pb = doc.add_paragraph(style='List Bullet')
                p_pb.paragraph_format.space_before = Pt(0)
                p_pb.paragraph_format.space_after = Pt(1.5)
                r_pb = p_pb.add_run(b)
                r_pb.font.size = Pt(9.5)

    # 7. Education Block
    education = resume_data.get('education')
    if education:
        add_section_header("Education")
        p_edu = doc.add_paragraph()
        p_edu.paragraph_format.space_before = Pt(2)
        p_edu.paragraph_format.space_after = Pt(2)

        if isinstance(education, dict):
            r_deg = p_edu.add_run(education.get('degree', 'Degree'))
            r_deg.bold = True
            r_deg.font.size = Pt(10)

            if education.get('year'):
                r_yr = p_edu.add_run(f"  |  {education.get('year')}")
                r_yr.font.size = Pt(9.5)
                r_yr.font.color.rgb = RGBColor(100, 116, 139)

            p_inst = doc.add_paragraph()
            p_inst.paragraph_format.space_before = Pt(0)
            p_inst.paragraph_format.space_after = Pt(2)
            r_in = p_inst.add_run(education.get('institution', 'University / College'))
            r_in.font.size = Pt(9.5)
            if education.get('gpa'):
                r_gpa = p_inst.add_run(f"  •  CGPA: {education.get('gpa')}")
                r_gpa.font.size = Pt(9.5)
        elif isinstance(education, list):
            for e in education:
                p_e = doc.add_paragraph(style='List Bullet')
                p_e.paragraph_format.space_before = Pt(0)
                p_e.paragraph_format.space_after = Pt(1.5)
                r_e = p_e.add_run(str(e))
                r_e.font.size = Pt(9.5)

    # 8. Certifications & Badges
    certifications = resume_data.get('certifications')
    if certifications:
        add_section_header("Certifications & Key Accreditations")
        for cert in certifications:
            p_c = doc.add_paragraph(style='List Bullet')
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(1.5)
            r_c = p_c.add_run(str(cert))
            r_c.font.size = Pt(9.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
